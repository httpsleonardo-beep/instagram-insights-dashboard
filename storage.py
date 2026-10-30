"""Histórico transacional e exportações derivadas do mesmo snapshot."""
from __future__ import annotations

import json
import os
import sqlite3
import tempfile
from datetime import datetime
from pathlib import Path

from filters import snapshot_selection
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

JSON_NAME = "insights_instagram_30dias.json"
EXCEL_NAME = "insights_instagram_30dias.xlsx"

def export_names(snapshot):
    selected = snapshot_selection(snapshot)
    if selected == {'scope': 'account', 'days': 30}:
        return EXCEL_NAME, JSON_NAME
    name = 'insights_' + selected['scope'] + ('_' + str(selected['days']) + 'dias' if 'days' in selected else '')
    return name + '.xlsx', name + '.json'


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class Store:
    def __init__(self, folder: Path):
        self.folder = folder
        folder.mkdir(parents=True, exist_ok=True)
        self.db = folder / "history.sqlite3"
        with self.connect() as connection:
            connection.execute("CREATE TABLE IF NOT EXISTS runs (id INTEGER PRIMARY KEY, timestamp TEXT NOT NULL, account_id TEXT, status TEXT NOT NULL, message TEXT, snapshot TEXT)")

    def connect(self):
        return sqlite3.connect(self.db, timeout=20)

    def save(self, snapshot):
        with self.connect() as connection:
            cursor = connection.execute("INSERT INTO runs(timestamp,account_id,status,message,snapshot) VALUES(?,?,?,?,?)", (snapshot["collected_at"], snapshot["account"]["id"], snapshot["status"], "; ".join(snapshot["warnings"]), json.dumps(snapshot, ensure_ascii=False)))
            return cursor.lastrowid

    def failure(self, timestamp, code, message, account_id=None):
        with self.connect() as connection:
            connection.execute("INSERT INTO runs(timestamp,account_id,status,message) VALUES(?,?,?,?)", (timestamp, account_id, code, message))

    def latest(self, account_id=None):
        with self.connect() as connection:
            row = connection.execute("SELECT snapshot FROM runs WHERE snapshot IS NOT NULL" + (" AND account_id=?" if account_id else "") + " ORDER BY id DESC LIMIT 1", (account_id,) if account_id else ()).fetchone()
        return json.loads(row[0]) if row else None

    def latest_for(self, account_id, selected):
        if not account_id:
            return None
        with self.connect() as connection:
            for row in connection.execute('SELECT snapshot FROM runs WHERE account_id=? AND snapshot IS NOT NULL ORDER BY id DESC', (account_id,)):
                snapshot = json.loads(row[0])
                if snapshot_selection(snapshot) == selected:
                    return snapshot
        return None

    def history(self, account_id=None, limit=None):
        sql = "SELECT id,timestamp,account_id,status,message FROM runs"
        params = []
        if account_id:
            sql += " WHERE account_id=?"
            params.append(account_id)
        sql += " ORDER BY id DESC"
        if limit:
            sql += " LIMIT ?"
            params.append(limit)
        with self.connect() as connection:
            connection.row_factory = sqlite3.Row
            return [dict(r) for r in connection.execute(sql, params)]

    def export(self, snapshot):
        history = self.history(snapshot["account"]["id"])
        excel_name, json_name = export_names(snapshot)
        atomic_json(self.folder / json_name, snapshot | {"collection_history": history})
        export_excel(self.folder / excel_name, snapshot, history)


def atomic_json(path, data):
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(data, stream, ensure_ascii=False, indent=2, allow_nan=False)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def build_sheet(book, name, title, headers, rows, widths, table_name, note):
    sheet = book.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.cell(1, 1, title).font = Font(name="Calibri", size=19, bold=True, color="FFFFFF")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor="122D32")
    sheet.row_dimensions[1].height = 37
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.cell(2, 1, note).font = Font(name="Calibri", size=10, color="52676C")
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 42
    for col, label in enumerate(headers, 1):
        sheet.cell(4, col, label)
    for row in rows:
        sheet.append([safe_cell(v) for v in row])
    for row in sheet.iter_rows(min_row=4):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="17373D")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0' if isinstance(cell.value, int) or cell.value.is_integer() else '#,##0.00'
        if row[0].row > 4:
            sheet.row_dimensions[row[0].row].height = 45
    for cell in sheet[4]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B62")
    sheet.row_dimensions[4].height = 25
    for index, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(index)].width = width
    if rows:
        from openpyxl.utils import get_column_letter
        table = Table(displayName=table_name, ref=f"A4:{get_column_letter(len(headers))}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    sheet.freeze_panes = "B5"
    sheet.print_title_rows = "1:4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def export_excel(path, snapshot, history):
    book = Workbook()
    book.remove(book.active)
    period = snapshot["period"]
    scope = f"Conta: {snapshot['account'].get('username', snapshot['account']['id'])} | {period['start'] or 'Data não informada'} a {period['end'] or 'Data não informada'} | Fuso: {period['timezone']} | Coleta: {snapshot['collected_at']}"
    general = []
    for item in snapshot["metrics"].values():
        source = item["source"] or {}
        method = {"network": "API do painel", "DOM": "Interface (aprox.)"}.get(source.get("method"), "Indisponível")
        general.append([item["label"], item["value"], "Atual na coleta" if item["label"] == "Total de seguidores" else period["label"], method, source.get("path", source.get("label", "")), item["definition"]])
    build_sheet(book, "Métricas Gerais", "Instagram · Métricas gerais", ["Métrica", "Valor", "Período", "Origem", "Campo da fonte", "Definição"], general, [32, 19, 24, 18, 55, 66], "MetricasGerais", scope)
    audience = snapshot.get('view_filters', {}).get('audience', 'all')
    views_field = 'views' if audience == 'all' else 'views_' + audience
    content = [[row["type"], row.get(views_field), row["interactions"], row["code"], "API do painel" if row["sources"] else "Indisponível"] for row in snapshot["content"]]
    scope += ' | Filtros de visualização: ' + str(snapshot.get('view_filters', {}))
    build_sheet(book, "Engajamento & Alcance", "Desempenho por conteúdo", ["Tipo de conteúdo", "Visualizações Instagram", "Interações", "Código original", "Origem"], content, [30, 29, 22, 22, 30], "Conteudo", scope + " | Célula vazia: indisponível. Alcance único por tipo não é retornado por este painel; não somar alcance de publicações.")
    # O alcance consolidado fica visível também na aba solicitada pelo usuário.
    sheet = book["Engajamento & Alcance"]
    index = sheet.max_row + 3
    for key in ("reach", "engaged_accounts"):
        if key not in snapshot["metrics"]:
            continue
        sheet.cell(index, 1, snapshot["metrics"][key]["label"])
        sheet.cell(index, 2, snapshot["metrics"][key]["value"])
        sheet.cell(index, 3, "Conta · " + period["label"])
        sheet.cell(index, 2).number_format = '#,##0'
        index += 1
    build_sheet(book, "Histórico de Coleta", "Histórico de execuções", ["Execução", "Data e hora", "Conta", "Status", "Observação"], [[r["id"], r["timestamp"], r["account_id"], r["status"], r["message"]] for r in history], [14, 31, 24, 25, 105], "Historico", "Histórico persistente desta conta. Falhas mantêm os últimos dados válidos. Horários incluem o deslocamento UTC.")
    selected = snapshot.get('selection', {})
    if selected.get('scope') == 'content':
        build_sheet(book, 'Conteúdo filtrado', 'Insights sobre o conteúdo', ['Publicação', 'Tipo', selected['metric']],
            [[r.get('shortcode') or r['id'], r['type'], r['value']] for r in snapshot.get('items', [])],
            [35, 25, 35], 'ConteudoFiltrado', scope + ' | ' + str(selected) + ' | ' + '; '.join(snapshot['warnings']))
    if selected.get('scope') == 'ads':
        build_sheet(book, 'Ferramentas para anúncios', 'Resultados acumulados por anúncio', ['Anúncio', 'Status', 'Objetivo', 'Resultado', 'Visualizações', 'Gasto', 'Orçamento', 'Público', 'Novos seguidores', 'Início', 'Término'],
            [[r['id'], r['status'], r['goal'], r['goal_value'], r['views'], r['spent'], r['budget'], r['audience'], r['followers'],
              datetime.fromtimestamp(r['start_epoch'], ZoneInfo('America/Sao_Paulo')).isoformat() if r.get('start_epoch') else None,
              datetime.fromtimestamp(r['end_epoch'], ZoneInfo('America/Sao_Paulo')).isoformat() if r.get('end_epoch') else None] for r in snapshot.get('items', [])],
            [28, 25, 25, 18, 20, 18, 18, 45, 22, 28, 28], 'Anuncios', scope + ' | Filtros: ' + str(snapshot.get('view_filters', {})) + ' | Valores acumulados; datas filtram anúncios, não o período dos resultados.')
    if snapshot.get('activity'):
        build_sheet(book, 'Horários ativos', 'Horários mais ativos dos seguidores', ['Dia', 'Faixa de horas', 'Seguidores'], [[r['day'],r['hours'],r['value']] for r in snapshot['activity'] if r['day'] == snapshot.get('view_filters', {}).get('weekday', r['day'])], [22,25,24], 'HorariosAtivos', 'Campo hourly_30_day: janela própria de 30 dias, independente do período da conta.')
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    os.close(fd)
    try:
        book.save(temporary)
        os.replace(temporary, path)
    finally:
        book.close()
        if os.path.exists(temporary):
            os.unlink(temporary)
